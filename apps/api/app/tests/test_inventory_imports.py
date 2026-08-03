import uuid
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models.inventory_import import InventoryImport
from app.models.inventory_item import InventoryItem
from app.models.inventory_movement import InventoryMovement


HEADERS = [
    "internal_code",
    "name",
    "category",
    "subcategory",
    "brand",
    "supplier",
    "unit",
    "purchase_price_ars",
    "purchase_tax_rate_percentage",
    "profit_margin_percentage",
    "sale_price_ars",
    "minimum_stock",
    "stock",
    "is_active",
    "notes",
]


def _headers(tenant) -> dict[str, str]:
    return {"X-Tenant-Id": str(tenant.id)}


def _row(**overrides) -> dict:
    data = {
        "internal_code": "",
        "name": "Amoxicilina 50mg",
        "category": "medication",
        "subcategory": "",
        "brand": "VetLab",
        "supplier": "Proveedor Uno",
        "unit": "tablet",
        "purchase_price_ars": "1000",
        "purchase_tax_rate_percentage": "21",
        "profit_margin_percentage": "35",
        "sale_price_ars": "",
        "minimum_stock": "3",
        "stock": "10",
        "is_active": "true",
        "notes": "Carga inicial",
    }
    data.update(overrides)
    return data


def _xlsx(rows: list[dict], *, sheet_name: str = "Inventario", headers: list[str] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    selected_headers = headers or HEADERS
    sheet.append(selected_headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in selected_headers])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _upload(client, tenant, content: bytes, *, mode: str = "initial_load", filename: str = "inventario.xlsx"):
    return client.post(
        "/api/v1/inventory/import/preview",
        headers=_headers(tenant),
        data={"mode": mode},
        files={
            "file": (
                filename,
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def _confirm(client, tenant, import_payload: dict, *, rows: list[dict] | None = None, reason: str = "Test import"):
    selected_rows = rows
    if selected_rows is None:
        selected_rows = [
            {
                "row_id": row["id"],
                "selected": row["status"] == "valid",
                "action": row["proposed_action"] if row["status"] == "valid" else "skip",
            }
            for row in import_payload["rows"]
        ]
    return client.post(
        f"/api/v1/inventory/import/{import_payload['id']}/confirm",
        headers=_headers(tenant),
        json={
            "explicit_confirm": True,
            "rows": selected_rows,
            "reason": reason,
        },
    )


def _create_item(db_session, tenant, **overrides) -> InventoryItem:
    item = InventoryItem(
        id=uuid.uuid4(),
        tenant_id=tenant.id,
        internal_code=overrides.pop("internal_code", "MED-00001"),
        name=overrides.pop("name", "Amoxicilina 50mg"),
        category=overrides.pop("category", "medication"),
        unit=overrides.pop("unit", "tablet"),
        brand=overrides.pop("brand", "VetLab"),
        supplier=overrides.pop("supplier", "Proveedor Uno"),
        current_stock=Decimal(str(overrides.pop("current_stock", "0"))),
        minimum_stock=Decimal(str(overrides.pop("minimum_stock", "3"))),
        purchase_price_ars=Decimal(str(overrides.pop("purchase_price_ars", "1000"))),
        purchase_tax_rate_percentage=Decimal(str(overrides.pop("purchase_tax_rate_percentage", "21"))),
        profit_margin_percentage=Decimal(str(overrides.pop("profit_margin_percentage", "35"))),
        sale_price_ars=Decimal(str(overrides.pop("sale_price_ars", "1633.50"))),
        sale_tax_rate_percentage=Decimal("0"),
        round_sale_price=False,
        is_active=overrides.pop("is_active", True),
        notes=overrides.pop("notes", None),
        **overrides,
    )
    db_session.add(item)
    db_session.commit()
    db_session.refresh(item)
    return item


def _inventory_item_count(db_session) -> int:
    return db_session.scalar(select(func.count()).select_from(InventoryItem))


def test_import_template_downloads_xlsx_with_expected_sheets_and_headers(client, tenant):
    response = client.get("/api/v1/inventory/import/template", headers=_headers(tenant))

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["Inventario", "Instrucciones", "Catálogos"]
    assert [cell.value for cell in next(workbook["Inventario"].iter_rows(max_row=1))] == HEADERS


def test_preview_rejects_non_xlsx_extension(client, tenant):
    response = _upload(client, tenant, b"not excel", filename="inventario.csv")

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_file_type"


def test_preview_rejects_corrupt_xlsx(client, tenant):
    response = _upload(client, tenant, b"not excel")

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_excel_file"


def test_preview_rejects_missing_inventory_sheet(client, tenant):
    response = _upload(client, tenant, _xlsx([_row()], sheet_name="Productos"))

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "missing_inventory_sheet"


def test_preview_rejects_missing_required_headers(client, tenant):
    response = _upload(client, tenant, _xlsx([_row()], headers=[header for header in HEADERS if header != "unit"]))

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "missing_headers"


def test_preview_rejects_row_limit(client, tenant):
    rows = [_row(name=f"Producto {index}") for index in range(2001)]

    response = _upload(client, tenant, _xlsx(rows))

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "too_many_rows"


def test_preview_reports_validation_errors_without_creating_products(client, tenant, db_session):
    before_count = _inventory_item_count(db_session)

    response = _upload(client, tenant, _xlsx([_row(name="", category="desconocida", unit="misterio")]))

    assert response.status_code == 201
    payload = response.json()["data"]
    assert payload["error_count"] == 1
    assert payload["rows"][0]["status"] == "error"
    assert {"name: required", "category: invalid", "unit: invalid"} <= set(payload["rows"][0]["errors"])
    assert _inventory_item_count(db_session) == before_count


def test_preview_normalizes_spanish_category_boolean_and_comma_decimals(client, tenant):
    response = _upload(
        client,
        tenant,
        _xlsx([_row(category="medicamento", purchase_price_ars="123,45", is_active="sí")]),
    )

    assert response.status_code == 201
    row = response.json()["data"]["rows"][0]
    assert row["normalized_data"]["category"] == "medication"
    assert row["normalized_data"]["purchase_price_ars"] == "123.45"
    assert row["normalized_data"]["is_active"] is True


def test_preview_rejects_ambiguous_decimal(client, tenant):
    response = _upload(client, tenant, _xlsx([_row(purchase_price_ars="1,234")]))

    assert response.status_code == 201
    row = response.json()["data"]["rows"][0]
    assert row["status"] == "error"
    assert "purchase_price_ars: ambiguous decimal value" in row["errors"]


def test_preview_matches_existing_code_only_inside_tenant(client, tenant, other_tenant, db_session):
    _create_item(db_session, other_tenant, internal_code="MED-00001", name="Producto de otra clínica")

    response = _upload(client, tenant, _xlsx([_row(internal_code="MED-00001")]))

    assert response.status_code == 201
    row = response.json()["data"]["rows"][0]
    assert row["status"] == "error"
    assert row["match_type"] == "conflict"
    assert "internal_code_not_found" in row["errors"]


def test_preview_exact_code_match_proposes_update(client, tenant, db_session):
    item = _create_item(db_session, tenant, internal_code="MED-00001", name="Amoxicilina 50mg")

    response = _upload(client, tenant, _xlsx([_row(internal_code=item.internal_code, name="Amoxicilina plus")]))

    assert response.status_code == 201
    row = response.json()["data"]["rows"][0]
    assert row["existing_inventory_item_id"] == str(item.id)
    assert row["match_type"] == "exact_match"
    assert row["proposed_action"] == "update"
    assert "name" in row["changed_fields"]


def test_preview_possible_duplicate_requires_review(client, tenant, db_session):
    _create_item(db_session, tenant, name="Amoxicilina 50mg", brand="VetLab", category="medication", unit="tablet")

    response = _upload(client, tenant, _xlsx([_row(internal_code="")]))

    assert response.status_code == 201
    row = response.json()["data"]["rows"][0]
    assert row["status"] == "warning"
    assert row["match_type"] == "possible_match"
    assert row["proposed_action"] == "review_required"
    assert "possible_duplicate" in row["warnings"]


def test_preview_marks_duplicate_rows_in_file(client, tenant):
    response = _upload(client, tenant, _xlsx([_row(name="Duplicado"), _row(name="Duplicado")]))

    assert response.status_code == 201
    rows = response.json()["data"]["rows"]
    assert [row["match_type"] for row in rows] == ["duplicate_in_file", "duplicate_in_file"]
    assert all(row["status"] == "error" for row in rows)


def test_catalog_update_confirmation_creates_product_without_stock_movement(client, tenant, db_session):
    response = _upload(client, tenant, _xlsx([_row(stock="99")]), mode="catalog_update")
    assert response.status_code == 201

    confirm_response = _confirm(client, tenant, response.json()["data"])

    assert confirm_response.status_code == 200
    item = db_session.scalar(select(InventoryItem).where(InventoryItem.tenant_id == tenant.id))
    assert item is not None
    assert item.current_stock == Decimal("0.00")
    assert db_session.scalar(select(func.count()).select_from(InventoryMovement)) == 0


def test_initial_load_confirmation_creates_product_and_initial_stock_movement(client, tenant, db_session):
    response = _upload(client, tenant, _xlsx([_row(stock="10")]))
    assert response.status_code == 201

    confirm_response = _confirm(client, tenant, response.json()["data"])

    assert confirm_response.status_code == 200
    payload = confirm_response.json()["data"]
    item = db_session.scalar(select(InventoryItem).where(InventoryItem.tenant_id == tenant.id))
    movement = db_session.scalar(select(InventoryMovement).where(InventoryMovement.tenant_id == tenant.id))
    assert item.current_stock == Decimal("10.00")
    assert movement.movement_type == "initial_stock"
    assert movement.quantity == Decimal("10.00")
    assert movement.source_type == "inventory_import"
    assert movement.source_id == payload["id"]
    assert movement.operation_id == uuid.UUID(payload["operation_id"])


def test_initial_load_update_creates_adjustment_in_for_positive_delta(client, tenant, db_session):
    item = _create_item(db_session, tenant, current_stock="4")
    response = _upload(client, tenant, _xlsx([_row(internal_code=item.internal_code, stock="9")]))
    assert response.status_code == 201

    confirm_response = _confirm(client, tenant, response.json()["data"])

    assert confirm_response.status_code == 200
    db_session.refresh(item)
    movement = db_session.scalar(select(InventoryMovement).where(InventoryMovement.tenant_id == tenant.id))
    assert item.current_stock == Decimal("9.00")
    assert movement.movement_type == "adjustment_in"
    assert movement.quantity == Decimal("5.00")


def test_initial_load_update_creates_adjustment_out_for_negative_delta(client, tenant, db_session):
    item = _create_item(db_session, tenant, current_stock="8")
    response = _upload(client, tenant, _xlsx([_row(internal_code=item.internal_code, stock="3")]))
    assert response.status_code == 201

    confirm_response = _confirm(client, tenant, response.json()["data"])

    assert confirm_response.status_code == 200
    db_session.refresh(item)
    movement = db_session.scalar(select(InventoryMovement).where(InventoryMovement.tenant_id == tenant.id))
    assert item.current_stock == Decimal("3.00")
    assert movement.movement_type == "adjustment_out"
    assert movement.quantity == Decimal("5.00")


def test_initial_load_empty_stock_updates_catalog_without_movement(client, tenant, db_session):
    item = _create_item(db_session, tenant, current_stock="8", notes="Antes")
    response = _upload(client, tenant, _xlsx([_row(internal_code=item.internal_code, stock="", notes="Después")]))
    assert response.status_code == 201

    confirm_response = _confirm(client, tenant, response.json()["data"])

    assert confirm_response.status_code == 200
    db_session.refresh(item)
    assert item.current_stock == Decimal("8.00")
    assert item.notes == "Después"
    assert db_session.scalar(select(func.count()).select_from(InventoryMovement)) == 0


def test_confirmation_skips_unselected_rows(client, tenant, db_session):
    response = _upload(client, tenant, _xlsx([_row(name="Crear"), _row(name="Omitir")]))
    payload = response.json()["data"]
    rows = [
        {"row_id": payload["rows"][0]["id"], "selected": True, "action": "create"},
        {"row_id": payload["rows"][1]["id"], "selected": False, "action": "skip"},
    ]

    confirm_response = _confirm(client, tenant, payload, rows=rows)

    assert confirm_response.status_code == 200
    assert _inventory_item_count(db_session) == 1
    assert confirm_response.json()["data"]["rows"][1]["status"] == "skipped"


def test_confirmation_blocks_double_confirm(client, tenant):
    response = _upload(client, tenant, _xlsx([_row(name="Producto")]))
    payload = response.json()["data"]
    assert _confirm(client, tenant, payload).status_code == 200

    second_response = _confirm(client, tenant, payload)

    assert second_response.status_code == 409
    assert second_response.json()["error"]["code"] == "inventory_import_already_confirmed"


def test_confirmation_blocks_same_hash_and_mode_after_confirmed_import(client, tenant):
    content = _xlsx([_row(name="Producto")])
    first_preview = _upload(client, tenant, content).json()["data"]
    assert _confirm(client, tenant, first_preview).status_code == 200
    second_preview = _upload(client, tenant, content)
    assert second_preview.status_code == 201

    duplicate_response = _confirm(client, tenant, second_preview.json()["data"])

    assert duplicate_response.status_code == 409
    assert duplicate_response.json()["error"]["code"] == "duplicate_import_file"


def test_confirmation_blocks_expired_preview(client, tenant, db_session):
    preview = _upload(client, tenant, _xlsx([_row(name="Producto")])).json()["data"]
    inventory_import = db_session.get(InventoryImport, uuid.UUID(preview["id"]))
    inventory_import.expires_at = datetime.now(UTC) - timedelta(minutes=1)
    db_session.add(inventory_import)
    db_session.commit()

    response = _confirm(client, tenant, preview)

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "inventory_import_expired"


def test_confirmation_detects_product_change_after_preview(client, tenant, db_session):
    item = _create_item(db_session, tenant)
    preview = _upload(client, tenant, _xlsx([_row(internal_code=item.internal_code, notes="Nuevo")])).json()["data"]
    item.notes = "Cambio concurrente"
    item.updated_at = item.updated_at + timedelta(seconds=5)
    db_session.add(item)
    db_session.commit()

    response = _confirm(client, tenant, preview)

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "inventory_item_changed"


def test_import_detail_confirm_and_history_are_tenant_scoped(client, tenant, other_tenant):
    preview = _upload(client, tenant, _xlsx([_row(name="Producto")])).json()["data"]

    detail_response = client.get(f"/api/v1/inventory/import/{preview['id']}", headers=_headers(other_tenant))
    confirm_response = _confirm(client, other_tenant, preview)
    history_response = client.get("/api/v1/inventory/imports", headers=_headers(other_tenant))

    assert detail_response.status_code == 404
    assert confirm_response.status_code == 404
    assert history_response.status_code == 200
    assert history_response.json()["data"] == []
