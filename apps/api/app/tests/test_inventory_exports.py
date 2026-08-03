import uuid
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models.inventory_item import InventoryItem
from app.models.inventory_movement import InventoryMovement
from app.repositories.inventory import InventoryRepository
from app.services.inventory_import import CANONICAL_HEADERS


def _headers(tenant) -> dict[str, str]:
    return {"X-Tenant-Id": str(tenant.id)}


def _create_item(db_session, tenant, **overrides) -> InventoryItem:
    item = InventoryItem(
        id=uuid.uuid4(),
        tenant_id=tenant.id,
        internal_code=overrides.pop("internal_code", f"MED-{uuid.uuid4().hex[:5]}"),
        name=overrides.pop("name", "Amoxicilina 50mg"),
        category=overrides.pop("category", "medication"),
        subcategory=overrides.pop("subcategory", None),
        brand=overrides.pop("brand", "VetLab"),
        supplier=overrides.pop("supplier", "Proveedor Uno"),
        unit=overrides.pop("unit", "tablet"),
        current_stock=Decimal(str(overrides.pop("current_stock", "0"))),
        minimum_stock=Decimal(str(overrides.pop("minimum_stock", "3"))),
        purchase_price_ars=overrides.pop("purchase_price_ars", Decimal("1000")),
        purchase_tax_rate_percentage=Decimal(str(overrides.pop("purchase_tax_rate_percentage", "21"))),
        profit_margin_percentage=Decimal(str(overrides.pop("profit_margin_percentage", "35"))),
        sale_price_ars=overrides.pop("sale_price_ars", Decimal("1633.50")),
        sale_tax_rate_percentage=Decimal("0"),
        round_sale_price=False,
        is_active=overrides.pop("is_active", True),
        notes=overrides.pop("notes", None),
        **overrides,
    )
    db_session.add(item)
    db_session.commit()
    db_session.refresh(item)
    return item


def _export(client, tenant, payload: dict):
    return client.post("/api/v1/inventory/export", headers=_headers(tenant), json=payload)


def _workbook(response):
    assert response.status_code == 200
    return load_workbook(BytesIO(response.content))


def _inventory_rows(response):
    sheet = _workbook(response)["Inventario"]
    rows = list(sheet.iter_rows(values_only=True))
    return rows[0], rows[1:]


def test_export_requires_tenant_context(client):
    response = client.post("/api/v1/inventory/export", json={"mode": "all"})

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "missing_tenant_header"


def test_export_rejects_invalid_mode_and_extra_fields(client, tenant):
    invalid_mode = _export(client, tenant, {"mode": "everything"})
    extra_field = _export(client, tenant, {"mode": "all", "unexpected": True})

    assert invalid_mode.status_code == 422
    assert extra_field.status_code == 422


def test_export_selected_requires_ids_and_limits_500(client, tenant):
    empty = _export(client, tenant, {"mode": "selected", "selected_ids": []})
    too_many = _export(
        client,
        tenant,
        {"mode": "selected", "selected_ids": [str(uuid.uuid4()) for _ in range(501)]},
    )

    assert empty.status_code == 422
    assert too_many.status_code == 422


def test_export_selected_rejects_other_tenant_id(client, tenant, other_tenant, db_session):
    other_item = _create_item(db_session, other_tenant)

    response = _export(
        client,
        tenant,
        {"mode": "selected", "selected_ids": [str(other_item.id)]},
    )

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "inventory_item_not_found"


def test_export_all_returns_xlsx_headers_sheets_filename_and_tenant_scope(
    client,
    tenant,
    other_tenant,
    db_session,
):
    own_item = _create_item(db_session, tenant, internal_code="MED-00001")
    _create_item(db_session, other_tenant, internal_code="MED-99999", name="Otro tenant")

    response = _export(client, tenant, {"mode": "all"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    disposition = response.headers["content-disposition"]
    assert "attachment;" in disposition
    assert "vetflow_inventario_" in disposition
    assert "/" not in disposition and "\\" not in disposition
    workbook = _workbook(response)
    assert workbook.sheetnames == ["Inventario", "Resumen", "Catálogos"]
    headers, rows = _inventory_rows(response)
    assert list(headers) == CANONICAL_HEADERS
    assert [row[0] for row in rows] == [own_item.internal_code]


def test_export_all_respects_is_active_filter(client, tenant, db_session):
    active = _create_item(db_session, tenant, internal_code="MED-00001", is_active=True)
    inactive = _create_item(db_session, tenant, internal_code="MED-00002", is_active=False)

    active_response = _export(client, tenant, {"mode": "all", "filters": {"is_active": True}})
    inactive_response = _export(client, tenant, {"mode": "all", "filters": {"is_active": False}})

    assert [row[0] for row in _inventory_rows(active_response)[1]] == [active.internal_code]
    assert [row[0] for row in _inventory_rows(inactive_response)[1]] == [inactive.internal_code]


def test_export_values_are_normalized_and_do_not_modify_database(client, tenant, db_session):
    before_movements = db_session.scalar(select(func.count()).select_from(InventoryMovement))
    _create_item(
        db_session,
        tenant,
        internal_code="MED-00001",
        name="=Formula peligrosa",
        subcategory=None,
        brand=None,
        supplier=None,
        current_stock="7.5",
        minimum_stock="2",
        purchase_price_ars=Decimal("12.30"),
        sale_price_ars=Decimal("20.50"),
        notes="+nota",
    )

    response = _export(client, tenant, {"mode": "all"})

    workbook = _workbook(response)
    sheet = workbook["Inventario"]
    assert sheet["B2"].value == "'=Formula peligrosa"
    assert sheet["B2"].data_type != "f"
    assert sheet["D2"].value is None
    assert sheet["E2"].value is None
    assert sheet["M2"].value == 7.5
    assert sheet["H2"].data_type == "n"
    assert sheet["J2"].data_type == "n"
    assert sheet["K2"].data_type == "n"
    assert sheet["O2"].value == "'+nota"
    assert db_session.scalar(select(func.count()).select_from(InventoryMovement)) == before_movements


def test_export_filtered_by_search_name_and_code(client, tenant, db_session):
    by_name = _create_item(db_session, tenant, internal_code="MED-00001", name="Meloxicam")
    by_code = _create_item(db_session, tenant, internal_code="ABC-12345", name="Otro")
    _create_item(db_session, tenant, internal_code="MED-00002", name="Sin match")

    response = _export(
        client,
        tenant,
        {"mode": "filtered", "filters": {"search": "melo", "sort_by": "name", "sort_direction": "asc"}},
    )
    code_response = _export(
        client,
        tenant,
        {"mode": "filtered", "filters": {"search": "abc-12345"}},
    )

    assert [row[0] for row in _inventory_rows(response)[1]] == [by_name.internal_code]
    assert [row[0] for row in _inventory_rows(code_response)[1]] == [by_code.internal_code]


def test_export_filtered_by_category_brand_supplier(client, tenant, db_session):
    matching = _create_item(
        db_session,
        tenant,
        internal_code="INS-00001",
        name="Jeringa",
        category="supply",
        brand="Acme",
        supplier="Central",
        unit="syringe",
    )
    _create_item(db_session, tenant, internal_code="MED-00001", category="medication", brand="Acme")

    response = _export(
        client,
        tenant,
        {
            "mode": "filtered",
            "filters": {
                "category": "supply",
                "brand": " acme ",
                "supplier": "central",
            },
        },
    )

    assert [row[0] for row in _inventory_rows(response)[1]] == [matching.internal_code]


def test_export_filtered_by_stock_statuses(client, tenant, db_session):
    in_stock = _create_item(db_session, tenant, internal_code="MED-00001", current_stock="10", minimum_stock="3")
    low_stock = _create_item(db_session, tenant, internal_code="MED-00002", current_stock="2", minimum_stock="3")
    out_of_stock = _create_item(db_session, tenant, internal_code="MED-00003", current_stock="0", minimum_stock="3")
    negative = _create_item(db_session, tenant, internal_code="MED-00004", current_stock="-1", minimum_stock="3")

    assert [row[0] for row in _inventory_rows(_export(client, tenant, {"mode": "filtered", "filters": {"stock_status": "in_stock"}}))[1]] == [in_stock.internal_code]
    assert [row[0] for row in _inventory_rows(_export(client, tenant, {"mode": "filtered", "filters": {"stock_status": "low_stock"}}))[1]] == [low_stock.internal_code]
    assert [row[0] for row in _inventory_rows(_export(client, tenant, {"mode": "filtered", "filters": {"stock_status": "out_of_stock"}}))[1]] == [out_of_stock.internal_code]
    assert [row[0] for row in _inventory_rows(_export(client, tenant, {"mode": "filtered", "filters": {"stock_status": "negative"}}))[1]] == [negative.internal_code]


def test_export_filtered_sorting_and_matches_list_without_pagination(client, tenant, db_session):
    _create_item(db_session, tenant, internal_code="MED-00001", name="B producto", brand="Marca")
    _create_item(db_session, tenant, internal_code="MED-00002", name="A producto", brand="Marca")
    _create_item(db_session, tenant, internal_code="MED-00003", name="C producto", brand="Marca")

    payload = {
        "mode": "filtered",
        "filters": {"brand": "Marca", "sort_by": "name", "sort_direction": "asc"},
    }
    export_response = _export(client, tenant, payload)
    list_response = client.get(
        "/api/v1/inventory/items",
        headers=_headers(tenant),
        params={"brand": "Marca", "sort_by": "name", "sort_direction": "asc", "page": 1, "page_size": 1},
    )
    desc_response = _export(
        client,
        tenant,
        {
            "mode": "filtered",
            "filters": {"brand": "Marca", "sort_by": "name", "sort_direction": "desc"},
        },
    )

    assert list_response.status_code == 200
    assert list_response.json()["meta"]["total"] == 3
    assert [row[1] for row in _inventory_rows(export_response)[1]] == [
        "A producto",
        "B producto",
        "C producto",
    ]
    assert [row[1] for row in _inventory_rows(desc_response)[1]] == [
        "C producto",
        "B producto",
        "A producto",
    ]


def test_export_selected_deduplicates_and_preserves_request_order(client, tenant, db_session):
    first = _create_item(db_session, tenant, internal_code="MED-00001", name="Primero")
    second = _create_item(db_session, tenant, internal_code="MED-00002", name="Segundo")

    response = _export(
        client,
        tenant,
        {
            "mode": "selected",
            "selected_ids": [str(second.id), str(first.id), str(second.id)],
        },
    )

    assert [row[0] for row in _inventory_rows(response)[1]] == [
        second.internal_code,
        first.internal_code,
    ]


def test_export_too_many_results_returns_clear_error(client, tenant, monkeypatch):
    def fake_list_items_for_export(*args, **kwargs):
        return [], 10_001

    monkeypatch.setattr(
        InventoryRepository,
        "list_items_for_export",
        fake_list_items_for_export,
    )

    response = _export(client, tenant, {"mode": "all"})

    assert response.status_code == 413
    assert response.json()["error"]["code"] == "inventory_export_too_large"


def test_export_zero_results_generates_valid_workbook(client, tenant):
    response = _export(
        client,
        tenant,
        {"mode": "filtered", "filters": {"search": "sin resultados"}},
    )

    workbook = _workbook(response)
    assert workbook["Inventario"].max_row == 1
    assert workbook["Resumen"]["B6"].value == 0


def test_exported_file_is_compatible_with_import_preview(client, tenant, db_session):
    item = _create_item(db_session, tenant, internal_code="MED-00001", name="Producto exportable", current_stock="8")
    export_response = _export(client, tenant, {"mode": "all"})
    assert export_response.status_code == 200

    catalog_preview = client.post(
        "/api/v1/inventory/import/preview",
        headers=_headers(tenant),
        data={"mode": "catalog_update"},
        files={
            "file": (
                "export.xlsx",
                export_response.content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    initial_preview = client.post(
        "/api/v1/inventory/import/preview",
        headers=_headers(tenant),
        data={"mode": "initial_load"},
        files={
            "file": (
                "export.xlsx",
                export_response.content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert catalog_preview.status_code == 201
    catalog_row = catalog_preview.json()["data"]["rows"][0]
    assert catalog_row["existing_inventory_item_id"] == str(item.id)
    assert catalog_row["match_type"] == "exact_match"
    assert catalog_row["errors"] == []
    assert catalog_row["expected_movement_type"] is None
    assert initial_preview.status_code == 201
    assert initial_preview.json()["data"]["rows"][0]["errors"] == []


def test_export_does_not_leave_xlsx_temporaries(client, tenant, db_session):
    before = set(Path("/tmp").glob("*.xlsx"))
    _create_item(db_session, tenant)

    response = _export(client, tenant, {"mode": "all"})

    assert response.status_code == 200
    assert set(Path("/tmp").glob("*.xlsx")) == before
