from __future__ import annotations

import logging
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.core.errors import AppError
from app.core.tenant import TenantContext
from app.models.inventory_item import InventoryItem
from app.models.tenant import Tenant
from app.repositories.inventory import InventoryRepository
from app.schemas.inventory import InventoryExportCreate, InventoryExportFilters
from app.services.inventory import (
    ALLOWED_SORT_BY,
    ALLOWED_SORT_ORDER,
    DEFAULT_SORT_BY,
    ZERO,
)
from app.services.inventory_import import CANONICAL_HEADERS, add_inventory_catalog_rows


logger = logging.getLogger(__name__)

MAX_EXPORT_ROWS = 10_000
MAX_SELECTED_EXPORT_IDS = 500
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_FILL = PatternFill("solid", fgColor="DDEAF7")
MONEY_FIELDS = {"purchase_price_ars", "sale_price_ars"}
DECIMAL_FIELDS = {
    "purchase_price_ars",
    "purchase_tax_rate_percentage",
    "profit_margin_percentage",
    "sale_price_ars",
    "minimum_stock",
    "stock",
}
TEXT_WIDTHS = {
    "internal_code": 16,
    "name": 32,
    "category": 16,
    "subcategory": 18,
    "brand": 20,
    "supplier": 24,
    "unit": 14,
    "notes": 36,
}


@dataclass(frozen=True)
class InventoryExportFile:
    content: bytes
    filename: str
    product_count: int


class InventoryExportService:
    def __init__(self, db: Session) -> None:
        self.db = db
        self.inventory_repository = InventoryRepository(db)

    def export_inventory(
        self,
        tenant: TenantContext,
        payload: InventoryExportCreate,
    ) -> InventoryExportFile:
        items, total, filters, sort_by, sort_direction = self._resolve_items(tenant.tenant_id, payload)
        if total > MAX_EXPORT_ROWS:
            raise AppError(
                413,
                "inventory_export_too_large",
                "La exportación supera 10.000 productos; aplica filtros para reducir el archivo.",
            )

        generated_at = datetime.now(UTC)
        workbook = self._build_workbook(
            tenant,
            items,
            payload=payload,
            filters=filters,
            sort_by=sort_by,
            sort_direction=sort_direction,
            generated_at=generated_at,
        )
        output = BytesIO()
        try:
            workbook.save(output)
        except Exception as exc:
            raise AppError(
                500,
                "inventory_export_failed",
                "No se pudo generar la exportación de inventario.",
            ) from exc

        logger.info(
            "inventory_export_generated",
            extra={
                "tenant_id": str(tenant.tenant_id),
                "user_id": str(tenant.user_id) if tenant.user_id else None,
                "mode": payload.mode,
                "product_count": total,
                "filters": self._filters_for_summary(filters),
            },
        )
        return InventoryExportFile(
            content=output.getvalue(),
            filename=self._build_filename(generated_at),
            product_count=total,
        )

    def _resolve_items(
        self,
        tenant_id: uuid.UUID,
        payload: InventoryExportCreate,
    ) -> tuple[list[InventoryItem], int, InventoryExportFilters | None, str, str]:
        if payload.mode == "selected":
            item_ids = self._dedupe_ids(payload.selected_ids)
            if len(item_ids) > MAX_SELECTED_EXPORT_IDS:
                raise AppError(
                    422,
                    "too_many_selected_ids",
                    "No se pueden exportar más de 500 productos seleccionados.",
                )
            items = self.inventory_repository.list_items_by_ids_for_export(tenant_id, item_ids)
            if len(items) != len(item_ids):
                raise AppError(
                    404,
                    "inventory_item_not_found",
                    "Uno o más productos seleccionados no existen.",
                )
            return items, len(items), None, "selected_ids", "request_order"

        filters = payload.filters or InventoryExportFilters()
        sort_by = filters.sort_by or DEFAULT_SORT_BY
        sort_direction = filters.sort_direction or "desc"
        if sort_by not in (ALLOWED_SORT_BY | {DEFAULT_SORT_BY}):
            raise AppError(422, "validation_error", "Invalid sort_by value")
        if sort_direction not in ALLOWED_SORT_ORDER:
            raise AppError(422, "validation_error", "Invalid sort_direction value")

        active_filter = filters.is_active
        if payload.mode == "all":
            filters = InventoryExportFilters(is_active=active_filter)

        items, total = self.inventory_repository.list_items_for_export(
            tenant_id,
            search=filters.search,
            category=filters.category,
            brand=filters.brand,
            supplier=filters.supplier,
            status=None,
            stock_status=filters.stock_status,
            is_active=filters.is_active,
            sort_by=sort_by,
            sort_direction=sort_direction,
            limit=MAX_EXPORT_ROWS + 1,
        )
        return items, total, filters, sort_by, sort_direction

    def _build_workbook(
        self,
        tenant: TenantContext,
        items: list[InventoryItem],
        *,
        payload: InventoryExportCreate,
        filters: InventoryExportFilters | None,
        sort_by: str,
        sort_direction: str,
        generated_at: datetime,
    ) -> Workbook:
        workbook = Workbook()
        inventory_sheet = workbook.active
        inventory_sheet.title = "Inventario"
        inventory_sheet.append(CANONICAL_HEADERS)
        self._style_header(inventory_sheet[1])
        inventory_sheet.freeze_panes = "A2"

        for item in items:
            inventory_sheet.append([self._export_value(header, item) for header in CANONICAL_HEADERS])

        inventory_sheet.auto_filter.ref = inventory_sheet.dimensions
        self._set_inventory_formats(inventory_sheet)

        summary_sheet = workbook.create_sheet("Resumen")
        self._write_summary_sheet(
            summary_sheet,
            tenant,
            payload=payload,
            filters=filters,
            sort_by=sort_by,
            sort_direction=sort_direction,
            generated_at=generated_at,
            product_count=len(items),
        )

        catalogs = workbook.create_sheet("Catálogos")
        add_inventory_catalog_rows(catalogs)
        self._set_catalog_widths(catalogs)
        return workbook

    def _write_summary_sheet(
        self,
        sheet,
        tenant: TenantContext,
        *,
        payload: InventoryExportCreate,
        filters: InventoryExportFilters | None,
        sort_by: str,
        sort_direction: str,
        generated_at: datetime,
        product_count: int,
    ) -> None:
        tenant_model = self.db.get(Tenant, tenant.tenant_id)
        tenant_name = tenant_model.display_name or tenant_model.name if tenant_model else ""
        requester = tenant.user_full_name or tenant.user_email or ""
        rows = [
            ("Exportación normalizada de inventario", ""),
            ("Fecha y hora", generated_at.isoformat()),
            ("Usuario solicitante", requester),
            ("Clínica", tenant_name),
            ("Modo", payload.mode),
            ("Productos", product_count),
            ("Filtros aplicados", self._filters_for_summary(filters)),
            ("Orden aplicado", f"{sort_by} {sort_direction}"),
            ("Advertencia", "El stock es una instantánea al momento de exportar."),
            (
                "Advertencia",
                "El archivo no modifica datos hasta importarse y confirmarse.",
            ),
            ("Advertencia", "internal_code no debe editarse."),
            ("Advertencia", "En catalog_update, stock se ignora."),
            ("Advertencia", "En initial_load, stock es valor objetivo."),
        ]
        for row in rows:
            sheet.append(row)
        sheet["A1"].font = Font(bold=True, size=14)
        for cell in sheet["A"]:
            cell.font = Font(bold=True)
        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 80

    def _export_value(self, header: str, item: InventoryItem) -> Any:
        value_by_header = {
            "internal_code": item.internal_code,
            "name": item.name,
            "category": item.category,
            "subcategory": item.subcategory,
            "brand": item.brand,
            "supplier": item.supplier,
            "unit": item.unit,
            "purchase_price_ars": item.purchase_price_ars,
            "purchase_tax_rate_percentage": item.purchase_tax_rate_percentage,
            "profit_margin_percentage": item.profit_margin_percentage,
            "sale_price_ars": item.sale_price_ars,
            "minimum_stock": item.minimum_stock,
            "stock": item.current_stock,
            "is_active": item.is_active,
            "notes": item.notes,
        }
        value = value_by_header[header]
        if value is None:
            return None
        if isinstance(value, str):
            return self._escape_excel_text(value)
        if isinstance(value, Decimal):
            return value
        return value

    def _set_inventory_formats(self, sheet) -> None:
        for index, header in enumerate(CANONICAL_HEADERS, start=1):
            column_letter = sheet.cell(row=1, column=index).column_letter
            sheet.column_dimensions[column_letter].width = TEXT_WIDTHS.get(header, 18)
            if header in DECIMAL_FIELDS:
                number_format = '#,##0.00' if header in MONEY_FIELDS else '0.00'
                for cell in sheet[column_letter][1:]:
                    cell.number_format = number_format

    def _set_catalog_widths(self, sheet) -> None:
        for column in ("A", "B", "C", "D"):
            sheet.column_dimensions[column].width = 28

    def _style_header(self, cells) -> None:
        for cell in cells:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL

    def _filters_for_summary(self, filters: InventoryExportFilters | None) -> str:
        if filters is None or not filters.has_any_filter():
            return "sin filtros"
        values = filters.model_dump(exclude_none=True)
        return ", ".join(f"{key}={value}" for key, value in values.items()) or "sin filtros"

    def _escape_excel_text(self, value: str) -> str:
        if value.startswith(("=", "+", "-", "@")):
            return f"'{value}"
        return value

    def _dedupe_ids(self, item_ids: list[uuid.UUID]) -> list[uuid.UUID]:
        unique_ids: list[uuid.UUID] = []
        seen: set[uuid.UUID] = set()
        for item_id in item_ids:
            if item_id in seen:
                continue
            seen.add(item_id)
            unique_ids.append(item_id)
        return unique_ids

    def _build_filename(self, generated_at: datetime) -> str:
        return f"vetflow_inventario_{generated_at.strftime('%Y-%m-%d_%H-%M')}.xlsx"
