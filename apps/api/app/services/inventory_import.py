from __future__ import annotations

import hashlib
import re
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import PurePath
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.core.errors import AppError
from app.models.inventory_import import InventoryImport, InventoryImportRow
from app.models.inventory_item import InventoryItem
from app.repositories.clinic import ClinicRepository
from app.repositories.inventory import InventoryRepository
from app.schemas.inventory import InventoryImportConfirmCreate
from app.services.inventory import INVENTORY_CATEGORY_PREFIXES, InventoryService, ZERO
from app.services.regional_settings import resolve_operational_money_settings

MAX_IMPORT_FILE_SIZE_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
MAX_IMPORT_COLUMNS = 40
PREVIEW_TTL_HOURS = 24
CANONICAL_HEADERS = [
    "internal_code",
    "name",
    "category",
    "subcategory",
    "brand",
    "supplier",
    "unit",
    "purchase_price_ars",
    "purchase_tax_rate_percentage",
    "profit_margin_percentage",
    "sale_price_ars",
    "minimum_stock",
    "stock",
    "is_active",
    "notes",
]
OPTIONAL_TEXT_FIELDS = {"internal_code", "subcategory", "brand", "supplier", "notes"}
DECIMAL_FIELDS = {
    "purchase_price_ars",
    "purchase_tax_rate_percentage",
    "profit_margin_percentage",
    "sale_price_ars",
    "minimum_stock",
    "stock",
}
IMPORT_MODES = {"initial_load", "catalog_update"}
UNITS = {
    "unit",
    "tablet",
    "capsule",
    "ampoule",
    "dose",
    "pipette",
    "bottle",
    "vial",
    "syringe",
    "ml",
    "liter",
    "gram",
    "kg",
    "pair",
    "box",
    "package",
    "other",
}
CATEGORY_ALIASES = {
    "medicamento": "medication",
    "medicamentos": "medication",
    "vacuna": "vaccine",
    "vacunas": "vaccine",
    "insumo": "supply",
    "insumos": "supply",
    "alimento": "food",
    "alimentos": "food",
    "accesorio": "accessory",
    "accesorios": "accessory",
    "otro": "other",
    "otros": "other",
}
BOOL_ALIASES = {
    "true": True,
    "false": False,
    "si": True,
    "sí": True,
    "no": False,
    "1": True,
    "0": False,
}


@dataclass
class ParsedRow:
    row_number: int
    values: dict[str, Any]
    json_values: dict[str, Any]
    errors: list[str]
    warnings: list[str]


class InventoryImportService:
    def __init__(self, db: Session) -> None:
        self.db = db
        self.inventory_repository = InventoryRepository(db)
        self.inventory_service = InventoryService(db)

    def build_template(self, tenant_id: uuid.UUID) -> bytes:
        settings = resolve_operational_money_settings(ClinicRepository(self.db), tenant_id)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Inventario"
        sheet.append(CANONICAL_HEADERS)
        sheet.append(
            [
                "",
                "Amoxicilina 50mg",
                "medication",
                "",
                "VetLab",
                "Proveedor Uno",
                "tablet",
                "1000",
                str(settings.default_purchase_tax_rate),
                str(settings.default_profit_margin),
                "",
                "3",
                "10",
                "true",
                "Ejemplo, borrar antes de importar",
            ]
        )
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEAF7")

        instructions = workbook.create_sheet("Instrucciones")
        for line in [
            "No cambies los encabezados de la hoja Inventario.",
            "Columnas obligatorias: name, category y unit. Para nuevos productos, Vetflow genera internal_code.",
            "Categorías: medication, vaccine, supply, food, accessory, other.",
            "Unidades: " + ", ".join(sorted(UNITS)) + ".",
            "Usa decimales con punto o coma. Evita valores ambiguos como 1,234.",
            f"IVA de compra vacío en nuevos productos usa el default de la clínica ({settings.default_purchase_tax_rate}%); vacío en actualizaciones conserva el valor actual.",
            "sale_price_ars es precio final al público; no se agrega IVA de venta.",
            "initial_load concilia el stock hasta el valor objetivo de stock.",
            "catalog_update ignora la columna stock y no modifica existencias.",
            f"Límite: {MAX_IMPORT_ROWS} filas y {MAX_IMPORT_FILE_SIZE_BYTES // (1024 * 1024)} MB.",
        ]:
            instructions.append([line])

        catalogs = workbook.create_sheet("Catálogos")
        add_inventory_catalog_rows(catalogs, settings.default_purchase_tax_rate)

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    async def preview_import(
        self,
        tenant_id: uuid.UUID,
        *,
        mode: str,
        upload_file,
        created_by_user_id: uuid.UUID | None,
    ) -> InventoryImport:
        if mode not in IMPORT_MODES:
            raise AppError(422, "validation_error", "Invalid import mode")

        filename = self._sanitize_filename(upload_file.filename)
        if not filename.lower().endswith(".xlsx"):
            raise AppError(422, "invalid_file_type", "Only .xlsx files are accepted")
        content = await upload_file.read(MAX_IMPORT_FILE_SIZE_BYTES + 1)
        if len(content) > MAX_IMPORT_FILE_SIZE_BYTES:
            raise AppError(413, "file_too_large", "Inventory import file is too large")
        file_hash = hashlib.sha256(content).hexdigest()

        parsed_rows = self._parse_workbook(content)
        if not parsed_rows:
            raise AppError(422, "empty_inventory_sheet", "Inventory sheet has no data rows")

        rows = self._classify_rows(tenant_id, mode, parsed_rows)
        summary = self._build_summary(rows)
        warnings = []
        if self.inventory_repository.has_confirmed_import_with_hash(
            tenant_id,
            file_hash=file_hash,
            mode=mode,
        ):
            warnings.append("Ya existe una importación confirmada con el mismo archivo y modo.")
        summary["warnings"] = warnings

        inventory_import = InventoryImport(
            tenant_id=tenant_id,
            created_by_user_id=created_by_user_id,
            mode=mode,
            status="preview",
            original_filename=filename,
            file_hash=file_hash,
            row_count=summary["row_count"],
            valid_count=summary["valid_count"],
            warning_count=summary["warning_count"],
            error_count=summary["error_count"],
            expires_at=datetime.now(UTC) + timedelta(hours=PREVIEW_TTL_HOURS),
            result_summary={"preview": self._json_ready(summary)},
        )
        self.inventory_repository.create_import(inventory_import)
        import_rows = [
            InventoryImportRow(
                tenant_id=tenant_id,
                import_id=inventory_import.id,
                **row,
            )
            for row in rows
        ]
        self.inventory_repository.create_import_rows(import_rows)
        self.db.commit()
        return self.get_import(tenant_id, inventory_import.id)

    def get_import(self, tenant_id: uuid.UUID, import_id: uuid.UUID) -> InventoryImport:
        inventory_import = self.inventory_repository.get_import_by_id(tenant_id, import_id)
        if inventory_import is None:
            raise AppError(404, "inventory_import_not_found", "Inventory import not found")
        self._mark_expired_if_needed(inventory_import)
        self._attach_summary(inventory_import)
        return inventory_import

    def list_imports(
        self,
        tenant_id: uuid.UUID,
        *,
        page: int,
        page_size: int,
    ) -> tuple[list[InventoryImport], dict]:
        if page < 1 or page_size < 1 or page_size > 100:
            raise AppError(422, "invalid_pagination", "Invalid pagination parameters")
        imports, total = self.inventory_repository.list_imports(
            tenant_id,
            page=page,
            page_size=page_size,
        )
        return imports, {
            "page": page,
            "page_size": page_size,
            "total": total,
            "total_pages": (total + page_size - 1) // page_size if total else 0,
        }

    def confirm_import(
        self,
        tenant_id: uuid.UUID,
        import_id: uuid.UUID,
        payload: InventoryImportConfirmCreate,
        *,
        created_by_user_id: uuid.UUID | None,
    ) -> InventoryImport:
        if not payload.explicit_confirm:
            raise AppError(422, "confirmation_required", "Explicit confirmation is required")

        inventory_import = self.inventory_repository.get_import_by_id(
            tenant_id,
            import_id,
            for_update=True,
        )
        if inventory_import is None:
            raise AppError(404, "inventory_import_not_found", "Inventory import not found")
        if inventory_import.status == "confirmed":
            raise AppError(409, "inventory_import_already_confirmed", "Inventory import already confirmed")
        if inventory_import.status != "preview":
            raise AppError(409, "inventory_import_not_confirmable", "Inventory import cannot be confirmed")
        if self._is_expired(inventory_import):
            inventory_import.status = "expired"
            self.db.commit()
            raise AppError(409, "inventory_import_expired", "Inventory import preview expired")
        if self.inventory_repository.has_confirmed_import_with_hash(
            tenant_id,
            file_hash=inventory_import.file_hash,
            mode=inventory_import.mode,
            exclude_import_id=inventory_import.id,
        ):
            raise AppError(409, "duplicate_import_file", "This file was already confirmed")

        actions = self._resolve_confirm_actions(inventory_import, payload)
        operation_id = uuid.uuid4()
        result = {
            "created_products": 0,
            "updated_products": 0,
            "skipped_rows": 0,
            "movements_created": 0,
            "movement_ids": [],
            "operation_id": str(operation_id),
        }

        try:
            for row in inventory_import.rows:
                action = actions.get(row.id, "skip")
                if action == "skip":
                    result["skipped_rows"] += 1
                    row.status = "skipped"
                    row.proposed_action = "skip"
                    continue
                if row.status == "error":
                    raise AppError(422, "invalid_import_row", f"Row {row.row_number} has validation errors")
                if row.proposed_action == "review_required":
                    raise AppError(422, "review_required", f"Row {row.row_number} requires review")

                data = self._row_data_with_decimals(row)
                item = self._create_or_update_item(
                    tenant_id,
                    row,
                    data,
                    action,
                    created_by_user_id=created_by_user_id,
                )
                if action == "create":
                    result["created_products"] += 1
                elif action == "update":
                    result["updated_products"] += 1

                movement = self._apply_stock_reconciliation(
                    tenant_id,
                    inventory_import,
                    row,
                    item,
                    data,
                    operation_id,
                    created_by_user_id,
                    payload.reason,
                )
                if movement is not None:
                    result["movements_created"] += 1
                    result["movement_ids"].append(str(movement.id))

            inventory_import.status = "confirmed"
            inventory_import.confirmed_at = datetime.now(UTC)
            inventory_import.operation_id = operation_id
            inventory_import.result_summary = {"result": result}
            self.db.add(inventory_import)
            self.db.commit()
        except AppError:
            self.db.rollback()
            raise
        except Exception as exc:
            self.db.rollback()
            raise AppError(500, "inventory_import_failed", "Inventory import failed") from exc

        return self.get_import(tenant_id, inventory_import.id)

    def _parse_workbook(self, content: bytes) -> list[ParsedRow]:
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
        except Exception as exc:
            raise AppError(422, "invalid_excel_file", "The uploaded file could not be read") from exc
        if "Inventario" not in workbook.sheetnames:
            raise AppError(422, "missing_inventory_sheet", "Workbook must include Inventario sheet")

        sheet = workbook["Inventario"]
        if sheet.max_column > MAX_IMPORT_COLUMNS:
            raise AppError(422, "too_many_columns", "Inventory sheet has too many columns")
        header_values = [
            self._normalize_header(cell.value)
            for cell in next(sheet.iter_rows(min_row=1, max_row=1))
        ]
        missing = [header for header in CANONICAL_HEADERS if header not in header_values]
        if missing:
            raise AppError(422, "missing_headers", f"Missing headers: {', '.join(missing)}")
        column_by_header = {
            header: header_values.index(header)
            for header in CANONICAL_HEADERS
        }

        parsed_rows: list[ParsedRow] = []
        for excel_row in sheet.iter_rows(min_row=2):
            if len(parsed_rows) >= MAX_IMPORT_ROWS:
                raise AppError(422, "too_many_rows", "Inventory import exceeds row limit")
            if all(cell.value in (None, "") for cell in excel_row):
                continue
            values: dict[str, Any] = {}
            json_values: dict[str, Any] = {}
            errors: list[str] = []
            warnings: list[str] = []
            for header in CANONICAL_HEADERS:
                cell = excel_row[column_by_header[header]]
                value = cell.value
                if cell.data_type == "f":
                    errors.append(f"{header}: formulas are not accepted")
                    value = None
                parsed_value, error = self._normalize_cell(header, value)
                if error:
                    errors.append(error)
                values[header] = parsed_value
                json_values[header] = self._json_value(parsed_value)
            self._validate_row_values(values, errors)
            parsed_rows.append(
                ParsedRow(
                    row_number=excel_row[0].row,
                    values=values,
                    json_values=json_values,
                    errors=errors,
                    warnings=warnings,
                )
            )
        return parsed_rows

    def _classify_rows(
        self,
        tenant_id: uuid.UUID,
        mode: str,
        parsed_rows: list[ParsedRow],
    ) -> list[dict]:
        internal_codes = {
            str(row.values["internal_code"]).lower()
            for row in parsed_rows
            if row.values.get("internal_code")
        }
        names = {
            self._normalize_lookup(row.values["name"])
            for row in parsed_rows
            if row.values.get("name")
        }
        candidates = self.inventory_repository.list_items_for_import_matching(
            tenant_id,
            internal_codes=internal_codes,
            names=names,
        )
        by_code = {item.internal_code.lower(): item for item in candidates}
        by_tuple: dict[tuple[str, str, str, str], list[InventoryItem]] = {}
        for item in candidates:
            key = self._item_match_key(
                item.name,
                item.category,
                item.brand,
                item.unit,
            )
            by_tuple.setdefault(key, []).append(item)

        duplicate_keys = self._duplicate_keys(parsed_rows)
        rows: list[dict] = []
        for parsed in parsed_rows:
            values = parsed.values
            errors = list(parsed.errors)
            warnings = list(parsed.warnings)
            duplicate_key = self._file_duplicate_key(values)
            existing_item = None
            match_type = "invalid" if errors else "new_product"
            proposed_action = "create"

            if duplicate_key in duplicate_keys:
                errors.append("duplicate_in_file")
                match_type = "duplicate_in_file"
                proposed_action = "review_required"

            code = values.get("internal_code")
            if not errors and code:
                existing_item = by_code.get(str(code).lower())
                if existing_item is None:
                    errors.append("internal_code_not_found")
                    match_type = "conflict"
                    proposed_action = "review_required"
                else:
                    match_type = "exact_match"
                    proposed_action = "update"
            elif not errors:
                possible_matches = by_tuple.get(
                    self._item_match_key(
                        values["name"],
                        values["category"],
                        values.get("brand"),
                        values["unit"],
                    ),
                    [],
                )
                if possible_matches:
                    existing_item = possible_matches[0] if len(possible_matches) == 1 else None
                    warnings.append("possible_duplicate")
                    match_type = "possible_match"
                    proposed_action = "review_required"

            status = "error" if errors else "warning" if warnings else "valid"
            snapshot = self._snapshot_item(existing_item) if existing_item is not None else None
            changed_fields = (
                self._changed_fields(existing_item, values)
                if existing_item is not None and proposed_action == "update"
                else []
            )
            stock_current, stock_target, stock_delta, expected_movement_type = self._stock_preview(
                mode,
                existing_item,
                values,
            )
            rows.append(
                {
                    "row_number": parsed.row_number,
                    "normalized_data": parsed.json_values,
                    "existing_inventory_item_id": existing_item.id if existing_item is not None else None,
                    "match_type": match_type,
                    "proposed_action": proposed_action,
                    "status": status,
                    "errors": errors,
                    "warnings": warnings,
                    "product_snapshot": snapshot,
                    "changed_fields": changed_fields,
                    "stock_current": stock_current,
                    "stock_target": stock_target,
                    "stock_delta": stock_delta,
                    "expected_movement_type": expected_movement_type,
                }
            )
        return rows

    def _create_or_update_item(
        self,
        tenant_id: uuid.UUID,
        row: InventoryImportRow,
        data: dict[str, Any],
        action: str,
        *,
        created_by_user_id: uuid.UUID | None,
    ) -> InventoryItem:
        if action == "create":
            item_data = self._item_create_data(tenant_id, data)
            item_data["internal_code"] = self.inventory_repository.get_next_internal_code(
                tenant_id,
                item_data["category"],
                INVENTORY_CATEGORY_PREFIXES[item_data["category"]],
            )
            item = InventoryItem(
                tenant_id=tenant_id,
                created_by_user_id=created_by_user_id,
                **item_data,
            )
            self.inventory_repository.create_item(item)
            return item

        if row.existing_inventory_item_id is None:
            raise AppError(422, "invalid_import_action", f"Row {row.row_number} has no product to update")
        item = self.inventory_repository.get_item_by_id_for_update(
            tenant_id,
            row.existing_inventory_item_id,
        )
        if item is None:
            raise AppError(404, "inventory_item_not_found", "Inventory item not found")
        snapshot = row.product_snapshot or {}
        if snapshot.get("updated_at") and snapshot["updated_at"] != item.updated_at.isoformat():
            raise AppError(409, "inventory_item_changed", f"Row {row.row_number} product changed after preview")
        updates = self._item_update_data(item, data)
        if updates:
            self.inventory_repository.update_item(item, updates)
        return item

    def _apply_stock_reconciliation(
        self,
        tenant_id: uuid.UUID,
        inventory_import: InventoryImport,
        row: InventoryImportRow,
        item: InventoryItem,
        data: dict[str, Any],
        operation_id: uuid.UUID,
        created_by_user_id: uuid.UUID | None,
        reason: str | None,
    ):
        if inventory_import.mode != "initial_load" or data.get("stock") is None:
            return None
        target = data["stock"]
        delta = target - item.current_stock
        if delta == ZERO:
            return None
        if delta > ZERO:
            movement_type = "initial_stock" if item.current_stock == ZERO else "adjustment_in"
            quantity = delta
        else:
            movement_type = "adjustment_out"
            quantity = abs(delta)
        return self.inventory_service.register_import_stock_movement(
            tenant_id,
            item.id,
            movement_type=movement_type,
            quantity=quantity,
            import_id=inventory_import.id,
            operation_id=operation_id,
            created_by_user_id=created_by_user_id,
            reason=reason or "Conciliación de inventario desde Excel",
            commit=False,
        )

    def _resolve_confirm_actions(
        self,
        inventory_import: InventoryImport,
        payload: InventoryImportConfirmCreate,
    ) -> dict[uuid.UUID, str]:
        row_by_id = {row.id: row for row in inventory_import.rows}
        actions: dict[uuid.UUID, str] = {}
        if not payload.rows:
            for row in inventory_import.rows:
                if row.status in {"valid", "warning"} and row.proposed_action in {"create", "update"}:
                    actions[row.id] = row.proposed_action
            return actions
        for selection in payload.rows:
            if selection.row_id not in row_by_id:
                raise AppError(422, "invalid_import_row", "Selected row does not belong to import")
            if not selection.selected or selection.action == "skip":
                actions[selection.row_id] = "skip"
                continue
            if selection.action not in {"create", "update"}:
                raise AppError(422, "invalid_import_action", "Invalid row action")
            row = row_by_id[selection.row_id]
            if selection.action == "create" and row.proposed_action != "create":
                raise AppError(422, "invalid_import_action", "Row cannot be created")
            if selection.action == "update" and row.proposed_action != "update":
                raise AppError(422, "invalid_import_action", "Row cannot be updated")
            actions[selection.row_id] = selection.action
        return actions

    def _item_create_data(
        self, tenant_id: uuid.UUID, data: dict[str, Any]
    ) -> dict[str, Any]:
        defaults = self.inventory_service._resolve_pricing_defaults(tenant_id)
        purchase_tax = data.get("purchase_tax_rate_percentage")
        if purchase_tax is None:
            purchase_tax = defaults["purchase_tax_rate"]
        profit_margin = data.get("profit_margin_percentage")
        if profit_margin is None:
            profit_margin = defaults["profit_margin"]
        item_data = {
            "name": data["name"],
            "category": data["category"],
            "subcategory": data.get("subcategory"),
            "brand": data.get("brand"),
            "supplier": data.get("supplier"),
            "unit": data["unit"],
            "minimum_stock": data.get("minimum_stock") or ZERO,
            "purchase_price_ars": data.get("purchase_price_ars"),
            "purchase_tax_rate_percentage": purchase_tax,
            "profit_margin_percentage": profit_margin,
            "sale_price_ars": data.get("sale_price_ars"),
            "sale_tax_rate_percentage": defaults["sale_tax_rate"],
            "round_sale_price": False,
            "notes": data.get("notes"),
            "is_active": True if data.get("is_active") is None else data["is_active"],
            "current_stock": ZERO,
        }
        item_data["sale_price_ars"] = self.inventory_service._resolve_sale_price(
            purchase_price_ars=item_data["purchase_price_ars"],
            purchase_tax_rate_percentage=item_data["purchase_tax_rate_percentage"],
            profit_margin_percentage=item_data["profit_margin_percentage"],
            round_sale_price=False,
            manual_sale_price_ars=data.get("sale_price_ars"),
            rounding_increment=defaults["rounding_increment"],
        )
        return item_data

    def _item_update_data(self, item: InventoryItem, data: dict[str, Any]) -> dict[str, Any]:
        updates: dict[str, Any] = {}
        for field in ("name", "category", "unit"):
            if data.get(field) is not None:
                updates[field] = data[field]
        for field in ("subcategory", "brand", "supplier", "notes"):
            if field in data:
                updates[field] = data[field]
        for field in (
            "minimum_stock",
            "purchase_price_ars",
            "purchase_tax_rate_percentage",
            "profit_margin_percentage",
            "sale_price_ars",
            "is_active",
        ):
            if data.get(field) is not None:
                updates[field] = data[field]
        if self.inventory_service._should_recalculate_sale_price(updates):
            updates["sale_price_ars"] = self.inventory_service._resolve_sale_price(
                purchase_price_ars=updates.get("purchase_price_ars", item.purchase_price_ars),
                purchase_tax_rate_percentage=updates.get(
                    "purchase_tax_rate_percentage",
                    item.purchase_tax_rate_percentage,
                ),
                profit_margin_percentage=updates.get(
                    "profit_margin_percentage",
                    item.profit_margin_percentage,
                ),
                round_sale_price=item.round_sale_price,
                manual_sale_price_ars=updates.get("sale_price_ars"),
                rounding_increment=self.inventory_service._resolve_pricing_defaults(
                    item.tenant_id
                )["rounding_increment"],
            )
        return updates

    def _row_data_with_decimals(self, row: InventoryImportRow) -> dict[str, Any]:
        data = dict(row.normalized_data)
        for field in DECIMAL_FIELDS:
            if data.get(field) is not None:
                data[field] = Decimal(str(data[field]))
        return data

    def _normalize_cell(self, header: str, value) -> tuple[Any, str | None]:
        if value is None:
            return None, None
        if header in DECIMAL_FIELDS:
            return self._parse_decimal(value, header)
        if header == "category":
            text = self._normalize_text(value)
            if text is None:
                return None, None
            lower = text.lower()
            return CATEGORY_ALIASES.get(lower, lower), None
        if header == "is_active":
            text = self._normalize_text(value)
            if text is None:
                return None, None
            normalized = text.lower()
            if normalized in BOOL_ALIASES:
                return BOOL_ALIASES[normalized], None
            return None, "is_active: invalid boolean value"
        text = self._normalize_text(value)
        return text, None

    def _validate_row_values(self, values: dict[str, Any], errors: list[str]) -> None:
        if not values.get("name"):
            errors.append("name: required")
        if values.get("name") and len(values["name"]) > 255:
            errors.append("name: too long")
        if values.get("category") not in INVENTORY_CATEGORY_PREFIXES:
            errors.append("category: invalid")
        if values.get("unit") not in UNITS:
            errors.append("unit: invalid")
        for field in ("purchase_tax_rate_percentage",):
            value = values.get(field)
            if value is not None and (value < ZERO or value > Decimal("100")):
                errors.append(f"{field}: must be between 0 and 100")
        for field in (
            "purchase_price_ars",
            "profit_margin_percentage",
            "sale_price_ars",
            "minimum_stock",
            "stock",
        ):
            value = values.get(field)
            if value is not None and value < ZERO:
                errors.append(f"{field}: cannot be negative")

    def _parse_decimal(self, value, field: str) -> tuple[Decimal | None, str | None]:
        if value in ("", None):
            return None, None
        if isinstance(value, int):
            return Decimal(value), None
        if isinstance(value, float):
            return Decimal(str(value)), None
        text = str(value).strip().replace(" ", "")
        if not text:
            return None, None
        if "," in text and "." in text:
            decimal_separator = "," if text.rfind(",") > text.rfind(".") else "."
            thousands_separator = "." if decimal_separator == "," else ","
            normalized = text.replace(thousands_separator, "").replace(decimal_separator, ".")
        elif "," in text:
            if self._looks_ambiguous_decimal(text, ","):
                return None, f"{field}: ambiguous decimal value"
            normalized = text.replace(",", ".")
        elif "." in text:
            if self._looks_ambiguous_decimal(text, "."):
                return None, f"{field}: ambiguous decimal value"
            normalized = text
        else:
            normalized = text
        try:
            return Decimal(normalized), None
        except InvalidOperation:
            return None, f"{field}: invalid decimal value"

    def _looks_ambiguous_decimal(self, value: str, separator: str) -> bool:
        parts = value.split(separator)
        return len(parts) == 2 and len(parts[1]) == 3 and len(parts[0]) <= 3

    def _stock_preview(
        self,
        mode: str,
        existing_item: InventoryItem | None,
        values: dict[str, Any],
    ) -> tuple[Decimal | None, Decimal | None, Decimal | None, str | None]:
        if mode != "initial_load" or values.get("stock") is None:
            return None, values.get("stock"), None, None
        current = existing_item.current_stock if existing_item is not None else ZERO
        target = values["stock"]
        delta = target - current
        if delta == ZERO:
            return current, target, delta, None
        if delta > ZERO:
            return current, target, delta, "initial_stock" if current == ZERO else "adjustment_in"
        return current, target, delta, "adjustment_out"

    def _changed_fields(self, item: InventoryItem, values: dict[str, Any]) -> list[str]:
        fields = [
            "name",
            "category",
            "subcategory",
            "brand",
            "supplier",
            "unit",
            "purchase_price_ars",
            "purchase_tax_rate_percentage",
            "profit_margin_percentage",
            "sale_price_ars",
            "minimum_stock",
            "is_active",
            "notes",
        ]
        changed = []
        for field in fields:
            value = values.get(field)
            if value is None and field in {"purchase_tax_rate_percentage", "profit_margin_percentage", "is_active"}:
                continue
            if str(getattr(item, field, None)) != str(value):
                changed.append(field)
        return changed

    def _snapshot_item(self, item: InventoryItem | None) -> dict | None:
        if item is None:
            return None
        return {
            "id": str(item.id),
            "internal_code": item.internal_code,
            "name": item.name,
            "category": item.category,
            "brand": item.brand,
            "unit": item.unit,
            "current_stock": str(item.current_stock),
            "updated_at": item.updated_at.isoformat(),
        }

    def _duplicate_keys(self, rows: list[ParsedRow]) -> set[str]:
        counts: dict[str, int] = {}
        for row in rows:
            key = self._file_duplicate_key(row.values)
            if key:
                counts[key] = counts.get(key, 0) + 1
        return {key for key, count in counts.items() if count > 1}

    def _file_duplicate_key(self, values: dict[str, Any]) -> str:
        if values.get("internal_code"):
            return f"code:{str(values['internal_code']).lower()}"
        if not values.get("name"):
            return ""
        return "tuple:" + "|".join(
            self._item_match_key(
                values.get("name"),
                values.get("category"),
                values.get("brand"),
                values.get("unit"),
            )
        )

    def _item_match_key(
        self,
        name: str | None,
        category: str | None,
        brand: str | None,
        unit: str | None,
    ) -> tuple[str, str, str, str]:
        return (
            self._normalize_lookup(name),
            category or "",
            self._normalize_lookup(brand),
            unit or "",
        )

    def _build_summary(self, rows: list[dict]) -> dict:
        stock_increase = ZERO
        stock_decrease = ZERO
        for row in rows:
            delta = row.get("stock_delta")
            if delta is None:
                continue
            if delta > ZERO:
                stock_increase += delta
            elif delta < ZERO:
                stock_decrease += abs(delta)
        return {
            "row_count": len(rows),
            "valid_count": sum(1 for row in rows if row["status"] == "valid"),
            "warning_count": sum(1 for row in rows if row["status"] == "warning"),
            "error_count": sum(1 for row in rows if row["status"] == "error"),
            "create_count": sum(1 for row in rows if row["proposed_action"] == "create"),
            "update_count": sum(1 for row in rows if row["proposed_action"] == "update"),
            "skip_count": sum(1 for row in rows if row["proposed_action"] == "skip"),
            "movement_count": sum(1 for row in rows if row.get("expected_movement_type")),
            "stock_increase_total": stock_increase,
            "stock_decrease_total": stock_decrease,
            "warnings": [],
        }

    def _attach_summary(self, inventory_import: InventoryImport) -> None:
        summary = self._build_summary(
            [
                {
                    "status": row.status,
                    "proposed_action": row.proposed_action,
                    "expected_movement_type": row.expected_movement_type,
                    "stock_delta": row.stock_delta,
                }
                for row in inventory_import.rows
            ]
        )
        if inventory_import.status == "confirmed" and inventory_import.result_summary:
            result = inventory_import.result_summary.get("result", {})
            summary["create_count"] = result.get("created_products", summary["create_count"])
            summary["update_count"] = result.get("updated_products", summary["update_count"])
            summary["skip_count"] = result.get("skipped_rows", summary["skip_count"])
            summary["movement_count"] = result.get("movements_created", summary["movement_count"])
        elif inventory_import.result_summary:
            summary["warnings"] = inventory_import.result_summary.get("preview", {}).get("warnings", [])
        setattr(inventory_import, "summary", summary)

    def _mark_expired_if_needed(self, inventory_import: InventoryImport) -> None:
        if inventory_import.status == "preview" and self._is_expired(inventory_import):
            inventory_import.status = "expired"
            self.db.add(inventory_import)
            self.db.commit()

    def _is_expired(self, inventory_import: InventoryImport) -> bool:
        expires_at = inventory_import.expires_at
        now = datetime.now(expires_at.tzinfo or UTC)
        if expires_at.tzinfo is None:
            now = now.replace(tzinfo=None)
        return expires_at <= now

    def _normalize_header(self, value) -> str:
        return str(value or "").strip().lower()

    def _normalize_text(self, value) -> str | None:
        text = str(value).strip() if value is not None else ""
        return text or None

    def _normalize_lookup(self, value: str | None) -> str:
        return re.sub(r"\s+", " ", (value or "").strip().lower())

    def _sanitize_filename(self, filename: str | None) -> str:
        name = PurePath(filename or "inventario.xlsx").name.strip() or "inventario.xlsx"
        return name[:255]

    def _json_value(self, value) -> Any:
        if isinstance(value, Decimal):
            return str(value)
        return value

    def _json_ready(self, value) -> Any:
        if isinstance(value, Decimal):
            return str(value)
        if isinstance(value, uuid.UUID):
            return str(value)
        if isinstance(value, list):
            return [self._json_ready(item) for item in value]
        if isinstance(value, dict):
            return {key: self._json_ready(item) for key, item in value.items()}
        return value


def add_inventory_catalog_rows(sheet, default_purchase_tax_rate: Decimal) -> None:
    sheet.append(["category", "unit", "boolean", "purchase_tax_rate_percentage"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEAF7")
    catalog_rows = max(len(INVENTORY_CATEGORY_PREFIXES), len(UNITS), len(BOOL_ALIASES), 3)
    tax_values: list[str] = []
    for value in (str(default_purchase_tax_rate), "0", "21", "personalizado"):
        if value not in tax_values:
            tax_values.append(value)
    for index in range(catalog_rows):
        sheet.append(
            [
                list(INVENTORY_CATEGORY_PREFIXES)[index]
                if index < len(INVENTORY_CATEGORY_PREFIXES)
                else "",
                sorted(UNITS)[index] if index < len(UNITS) else "",
                list(BOOL_ALIASES)[index] if index < len(BOOL_ALIASES) else "",
                tax_values[index] if index < len(tax_values) else "",
            ]
        )
